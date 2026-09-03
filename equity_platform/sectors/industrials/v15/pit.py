from __future__ import annotations

from hashlib import sha256
import json
from pathlib import Path
import re

import numpy as np
from openpyxl import load_workbook
import pandas as pd


VINTAGE_COLUMNS = [
    "source", "dataset", "series_id", "observation_period", "vintage_value",
    "release_date", "available_at", "revision_number", "source_hash",
    "vintage_status", "pit_eligible", "report_period", "source_path",
    "value_changed_from_prior_vintage",
]


def build_forecast_origins(ir_history: pd.DataFrame) -> pd.DataFrame:
    """Create one-quarter-ahead forecast cutoffs from CAT's public IR releases.

    A target quarter can only use information available at the immediately
    preceding quarterly earnings release.  The target's own filing date is
    retained as the settlement timestamp and is never admitted to features.
    """
    releases = (
        ir_history[["period", "filing_date"]]
        .drop_duplicates()
        .assign(
            quarter=lambda frame: pd.PeriodIndex(frame["period"], freq="Q"),
            filing_date=lambda frame: pd.to_datetime(frame["filing_date"]),
        )
        .sort_values("quarter")
        .reset_index(drop=True)
    )
    rows: list[dict[str, object]] = []
    for index in range(1, len(releases)):
        previous = releases.iloc[index - 1]
        target = releases.iloc[index]
        if target["quarter"].ordinal - previous["quarter"].ordinal != 1:
            continue
        rows.append(
            {
                "period": str(target["quarter"]),
                "forecast_as_of": previous["filing_date"].date().isoformat(),
                "actual_available_at": target["filing_date"].date().isoformat(),
                "origin_period": str(previous["quarter"]),
                "horizon_quarters": 1,
                "actual_after_forecast": bool(target["filing_date"] > previous["filing_date"]),
            }
        )
    return pd.DataFrame(rows)


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _code(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().replace(".0", "")


def _release_date(workbook) -> pd.Timestamp:
    pattern = re.compile(r"ppi_(\d{2})(\d{2})(\d{4})[.\s]+htm", re.IGNORECASE)
    for row in workbook["Table of Contents"].iter_rows(values_only=True):
        for value in row:
            match = pattern.search(str(value))
            if match:
                month, day, year = (int(item) for item in match.groups())
                return pd.Timestamp(year, month, day)
    raise ValueError("PPI workbook does not expose its archived news-release date")


def _table_rows(workbook, sheet_name: str, targets: set[str], report_period: pd.Period) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    periods = pd.period_range(end=report_period, periods=5, freq="M")
    for row in workbook[sheet_name].iter_rows(min_row=7, values_only=True):
        values = list(row)
        if len(values) < 10:
            continue
        if sheet_name == "Table 9":
            group, item = _code(values[2]), _code(values[3])
            series_id = f"WPU{group}{item}" if group and item else ""
        else:
            industry, product = _code(values[2]), _code(values[3])
            series_id = f"PCU{industry}{industry}" if industry and not product else ""
        if series_id not in targets:
            continue
        for period, raw_value in zip(periods, values[5:10]):
            value = pd.to_numeric(raw_value, errors="coerce")
            if pd.notna(value):
                rows.append(
                    {
                        "series_id": series_id,
                        "observation_period": str(period),
                        "vintage_value": float(value),
                    }
                )
    return rows


def parse_bls_ppi_vintages(
    *,
    project_root: Path,
    archive_manifest_path: Path,
    sensor_map_path: Path,
    cutoff: pd.Timestamp,
) -> dict[str, pd.DataFrame]:
    manifest = json.loads(archive_manifest_path.read_text(encoding="utf-8"))
    sensors = pd.read_csv(sensor_map_path)
    targets = set(sensors["series_id"])
    parsed: list[dict[str, object]] = []
    source_rows: list[dict[str, object]] = []
    schedule_rows: list[dict[str, object]] = []
    cutoff_date = pd.Timestamp(cutoff).date()
    for artifact in manifest.get("release_schedule_artifacts", []):
        path = project_root / artifact["local_path"]
        actual_hash = _sha256(path)
        if actual_hash != artifact["sha256"]:
            raise ValueError(f"BLS release-schedule hash mismatch: {path}")
        schedule_rows.append(
            {
                "calendar_year": artifact["calendar_year"],
                "source_path": str(path),
                "source_url": artifact["source_url"],
                "source_sha256": actual_hash,
                "hash_match": True,
                "supporting_release_calendar_only": True,
                "pdf_parsing_used": False,
            }
        )
    for artifact in manifest["ppi_report_artifacts"]:
        path = project_root / artifact["local_path"]
        actual_hash = _sha256(path)
        if actual_hash != artifact["sha256"]:
            raise ValueError(f"BLS archived report hash mismatch: {path}")
        report_period = pd.Period(artifact["report_period"], freq="M")
        workbook = load_workbook(path, read_only=True, data_only=True)
        release_date = _release_date(workbook)
        if release_date.date() > cutoff_date:
            workbook.close()
            continue
        available_at = release_date.tz_localize("America/New_York") + pd.Timedelta(hours=8, minutes=30)
        available_at = available_at.tz_convert("UTC")
        report_rows = _table_rows(workbook, "Table 9", targets, report_period)
        report_rows.extend(_table_rows(workbook, "Table 11", targets, report_period))
        workbook.close()
        for row in report_rows:
            parsed.append(
                {
                    "source": "BLS_PPI_ARCHIVED_DETAILED_REPORT_XLSX",
                    "dataset": "PPI_AS_RELEASED_VINTAGE",
                    **row,
                    "release_date": release_date.date().isoformat(),
                    "available_at": available_at.isoformat(),
                    "source_hash": actual_hash,
                    "pit_eligible": True,
                    "report_period": str(report_period),
                    "source_path": str(path),
                }
            )
        source_rows.append(
            {
                "report_period": str(report_period),
                "release_date": release_date.date().isoformat(),
                "source_path": str(path),
                "source_url": artifact["source_url"],
                "source_sha256": actual_hash,
                "target_series_rows": len(report_rows),
                "hash_match": True,
                "pdf_parsing_used": False,
            }
        )
    frame = pd.DataFrame(parsed).drop_duplicates(["series_id", "observation_period", "release_date"]).sort_values(
        ["series_id", "observation_period", "release_date"]
    ).reset_index(drop=True)
    first_report = frame.groupby(["series_id", "observation_period"])["report_period"].transform("min")
    frame["revision_number"] = frame.groupby(["series_id", "observation_period"]).cumcount().astype("Int64")
    left_censored = pd.PeriodIndex(first_report, freq="M") > pd.PeriodIndex(frame["observation_period"], freq="M")
    frame.loc[left_censored, "revision_number"] = pd.NA
    prior = frame.groupby(["series_id", "observation_period"])["vintage_value"].shift(1)
    frame["value_changed_from_prior_vintage"] = prior.notna() & ~np.isclose(frame["vintage_value"], prior, equal_nan=False)
    frame["vintage_status"] = np.where(
        left_censored,
        "LEFT_CENSORED_KNOWN_AS_OF_RELEASE",
        np.where(frame["revision_number"].eq(0).fillna(False), "INITIAL_RELEASE", "SUBSEQUENT_PUBLICATION"),
    )
    frame = frame[VINTAGE_COLUMNS]
    missing_series = sorted(targets - set(frame["series_id"]))
    coverage = (
        frame.groupby("series_id")
        .agg(
            observation_months=("observation_period", "nunique"),
            published_vintages=("vintage_value", "size"),
            first_release_date=("release_date", "min"),
            last_release_date=("release_date", "max"),
            changed_republications=("value_changed_from_prior_vintage", "sum"),
        )
        .reset_index()
    )
    audit = pd.DataFrame(
        [
            {
                "archive_report_files": len(source_rows),
                "archive_report_hash_mismatches": 0,
                "release_schedule_files": len(schedule_rows),
                "release_schedule_hash_mismatches": 0,
                "required_series": len(targets),
                "parsed_series": frame["series_id"].nunique(),
                "missing_series": "|".join(missing_series),
                "vintage_rows": len(frame),
                "release_date_coverage_pct": float(frame["release_date"].notna().mean() * 100.0),
                "available_at_coverage_pct": float(frame["available_at"].notna().mean() * 100.0),
                "source_hash_coverage_pct": float(frame["source_hash"].notna().mean() * 100.0),
                "pit_eligible_rows": int(frame["pit_eligible"].sum()),
                "historical_pit_ready": not missing_series and bool(frame["pit_eligible"].all()),
                "pdf_parsing_used": False,
                "status": "HISTORICAL_PIT_ARCHIVE_READY" if not missing_series else "FAIL_CLOSED_MISSING_SERIES",
            }
        ]
    )
    return {
        "bls_ppi_archive_sources": pd.DataFrame(source_rows),
        "bls_release_schedule_sources": pd.DataFrame(schedule_rows),
        "bls_ppi_vintage_canonical": frame,
        "bls_ppi_vintage_coverage": coverage,
        "bls_ppi_vintage_audit": audit,
    }


def _quarter_values(selected: pd.DataFrame, series_id: str, quarter: pd.Period) -> pd.Series:
    months = {str(period) for period in pd.period_range(quarter.start_time, quarter.end_time, freq="M")}
    return selected.loc[
        selected["series_id"].eq(series_id) & selected["observation_period"].isin(months),
        "vintage_value",
    ]


def build_pit_forecast_features(
    *,
    vintages: pd.DataFrame,
    sensor_map: pd.DataFrame,
    forecast_origins: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    vintages = vintages.copy()
    vintages["release_date"] = pd.to_datetime(vintages["release_date"])
    rows: list[dict[str, object]] = []
    selection_rows: list[dict[str, object]] = []
    for _, origin in forecast_origins.iterrows():
        target = pd.Period(origin["period"], freq="Q")
        as_of = pd.Timestamp(origin["forecast_as_of"])
        available = vintages.loc[vintages["release_date"].le(as_of)].copy()
        selected = available.loc[
            available.groupby(["series_id", "observation_period"])["release_date"].idxmax()
        ]
        reference_quarter = None
        for lag in range(1, 5):
            candidate = target - lag
            prior = candidate - 4
            complete = True
            for series_id in sensor_map["series_id"].unique():
                if len(_quarter_values(selected, series_id, candidate)) != 3 or len(_quarter_values(selected, series_id, prior)) != 3:
                    complete = False
                    break
            if complete:
                reference_quarter = candidate
                break
        if reference_quarter is None:
            continue
        series_yoy: dict[str, float] = {}
        for series_id in sensor_map["series_id"].unique():
            current = _quarter_values(selected, series_id, reference_quarter)
            prior = _quarter_values(selected, series_id, reference_quarter - 4)
            series_yoy[series_id] = float(current.mean() / prior.mean() * 100.0 - 100.0)
            chosen = selected.loc[
                selected["series_id"].eq(series_id)
                & selected["observation_period"].isin(
                    {str(period) for period in pd.period_range((reference_quarter - 4).start_time, reference_quarter.end_time, freq="M")}
                )
            ]
            selection_rows.append(
                {
                    "period": str(target),
                    "forecast_as_of": as_of.date().isoformat(),
                    "feature_reference_quarter": str(reference_quarter),
                    "series_id": series_id,
                    "selected_vintage_rows": len(chosen),
                    "latest_selected_release_date": chosen["release_date"].max().date().isoformat(),
                    "cutoff_respected": bool(chosen["release_date"].le(as_of).all()),
                }
            )
        for segment, group in sensor_map.groupby("segment"):
            output = group.loc[group["role"].eq("output_price")]
            costs = group.loc[group["role"].ne("output_price")]
            rows.append(
                {
                    "period": str(target),
                    "segment": segment,
                    "forecast_as_of": as_of.date().isoformat(),
                    "actual_available_at": origin["actual_available_at"],
                    "feature_reference_quarter": str(reference_quarter),
                    "feature_lag_quarters": target.ordinal - reference_quarter.ordinal,
                    "output_price_yoy_pct": float(sum(row.weight * series_yoy[row.series_id] for row in output.itertuples())),
                    "dedicated_cost_yoy_pct": float(sum(row.weight * series_yoy[row.series_id] for row in costs.itertuples())),
                    "ppi_series_coverage_pct": 100.0,
                    "historical_pit_eligible": True,
                }
            )
    features = pd.DataFrame(rows).sort_values(["period", "segment"]).reset_index(drop=True)
    selection = pd.DataFrame(selection_rows).drop_duplicates().sort_values(["period", "series_id"]).reset_index(drop=True)
    summary = pd.DataFrame(
        [
            {
                "forecast_periods": features["period"].nunique(),
                "segments": features["segment"].nunique(),
                "feature_rows": len(features),
                "minimum_series_coverage_pct": float(features["ppi_series_coverage_pct"].min()),
                "maximum_feature_lag_quarters": int(features["feature_lag_quarters"].max()),
                "cutoff_violations": int((~selection["cutoff_respected"]).sum()),
                "historical_pit_ready": bool(features["historical_pit_eligible"].all() and selection["cutoff_respected"].all()),
            }
        ]
    )
    return {"pit_segment_features": features, "pit_vintage_selection_audit": selection, "pit_feature_summary": summary}
