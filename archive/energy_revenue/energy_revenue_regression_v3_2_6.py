# ============================================================
# ENERGY REVENUE NOWCAST V3.2.6
# ============================================================
#
# 핵심
# ------------------------------------------------------------
# COP / FANG / DVN:
#   기존 V3.2.3 structural panel 재사용
#
# EOG:
#   공식 EOG Investor Relations
#
#       Statistics - XLS
#           ↓
#       Actual Oil / NGL / Gas / Total Production
#
#       Guidance PDF
#           ↓
#       Next-quarter production guidance
#
#
# EOG BOE QC
# ------------------------------------------------------------
#
# Reconstructed Total MBoed
#   =
#   Oil MBod
#   + NGL MBbld
#   + Gas MMcfd / 6
#
#
# reported vs reconstructed 오차 <= 5%
#     -> reported total 사용
#
# 오차 > 5%
#     -> reconstructed total 사용
#        + quality penalty
#
#
# Structural
# ------------------------------------------------------------
#
# Revenue Log YoY
# ≈
# Price Mix Log YoY
# +
# Company Production Log YoY
#
#
# Blend
# ------------------------------------------------------------
#
# Final =
#     w * Structural
#     +
#     (1-w) * V2.1
#
#
# Metrics
# ------------------------------------------------------------
#
# 1. log-points MAE
# 2. 일반 Revenue YoY %p MAE
#
#
# REQUIRED FILES
# ------------------------------------------------------------
#
# ./data-lake/energy_v2_1_panel.csv
# ./data-lake/energy_v2_1_validation.csv
# ./data-lake/energy_v2_1_nowcast.csv
#
# 그리고 아래 중 하나:
#
# ./data-lake/energy_v3_2_3_structural_panel.csv
# ./data-lake/energy_v3_2_5_structural_panel.csv
# ./data-lake/energy_v3_2_4_structural_panel.csv
#
# ============================================================


import os
import re
import json
import time
import warnings

from io import BytesIO
from pathlib import Path
from datetime import date
from urllib.parse import urljoin

import numpy as np
import pandas as pd
import requests

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader


warnings.filterwarnings("ignore")


# ============================================================
# 0. CONFIG
# ============================================================

DATA_LAKE_DIR = Path(
    os.getenv(
        "DATA_LAKE_DIR",
        "./data-lake",
    )
)


TICKERS = [
    x.strip().upper()

    for x in os.getenv(
        "TICKERS",
        "COP,EOG,FANG,DVN",
    ).split(",")

    if x.strip()
]


AS_OF_DATE_ENV = os.getenv(
    "AS_OF_DATE"
)


REFRESH_EOG_IR = (
    os.getenv(
        "REFRESH_EOG_IR",
        "1",
    )
    .strip()
    .lower()

    in {
        "1",
        "true",
        "yes",
        "y",
    }
)


EOG_IR_START_YEAR = int(
    os.getenv(
        "EOG_IR_START_YEAR",
        "2022",
    )
)


EOG_EVENT_ITEM_MAX = int(
    os.getenv(
        "EOG_EVENT_ITEM_MAX",
        "80",
    )
)


EOG_MIN_EVENTS = int(
    os.getenv(
        "EOG_MIN_EVENTS",
        "14",
    )
)


EOG_RECON_ERROR_MAX_PCT = float(
    os.getenv(
        "EOG_RECON_ERROR_MAX_PCT",
        "5.0",
    )
)


HTTP_TIMEOUT = int(
    os.getenv(
        "HTTP_TIMEOUT",
        "40",
    )
)


# ============================================================
# BLEND CONFIG
# ============================================================

BLEND_WEIGHTS = np.round(
    np.arange(
        0.0,
        1.0001,
        0.05,
    ),
    2,
)


MIN_BLEND_HISTORY = int(
    os.getenv(
        "MIN_BLEND_HISTORY",
        "3",
    )
)


NO_COMPANY_PROD_MAX_WEIGHT = 0.25


# ============================================================
# SOURCE QUALITY
# ============================================================

SOURCE_QUALITY = {

    "GUIDANCE_VS_ACTUAL":
        1.00,

    "GUIDANCE_VS_GUIDANCE":
        0.80,

    "ACTUAL_VS_ACTUAL":
        0.90,

    "INDUSTRY_FALLBACK":
        0.35,
}


MIN_LOG_GROWTH = -100.0
MAX_LOG_GROWTH = 150.0


# ============================================================
# URLS
# ============================================================

EOG_BASE_URL = (
    "https://investors.eogresources.com"
)


EOG_EVENTS_URL = (
    EOG_BASE_URL
    + "/events-and-presentations"
)


EOG_INVESTORS_URL = (
    EOG_BASE_URL
    + "/"
)


# ============================================================
# 1. PATH / DATE
# ============================================================

DATA_LAKE_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


def lake_path(
    filename
):

    return (
        DATA_LAKE_DIR
        / filename
    )


def resolve_as_of():

    if AS_OF_DATE_ENV:

        return (
            pd.Timestamp(
                AS_OF_DATE_ENV
            )
            .normalize()
        )

    return pd.Timestamp(
        date.today()
    )


AS_OF = resolve_as_of()


print(
    "=" * 90
)

print(
    "ENERGY REVENUE NOWCAST V3.2.6"
)

print(
    "=" * 90
)

print(
    "AS OF          :",
    AS_OF.date()
)

print(
    "TICKERS        :",
    TICKERS
)

print(
    "DATA LAKE      :",
    DATA_LAKE_DIR.resolve()
)

print(
    "REFRESH EOG IR :",
    REFRESH_EOG_IR
)

print(
    "EOG START YEAR :",
    EOG_IR_START_YEAR
)


# ============================================================
# 2. HTTP SESSION
# ============================================================

session = requests.Session()


session.headers.update(
    {
        "User-Agent":
            (
                "Mozilla/5.0 "
                "(Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 "
                "Chrome/140 Safari/537.36"
            ),

        "Accept-Language":
            "en-US,en;q=0.9",
    }
)


def http_get(
    url,
):

    response = session.get(
        url,
        timeout=HTTP_TIMEOUT,
    )

    response.raise_for_status()

    return response


# ============================================================
# 3. GENERAL HELPERS
# ============================================================

def normalize_text(
    value
):

    if value is None:

        return ""

    value = str(value)

    value = (
        value
        .replace("\xa0", " ")
        .replace("−", "-")
        .replace("–", "-")
        .replace("—", "-")
    )

    value = re.sub(
        r"\s+",
        " ",
        value,
    )

    return value.strip()


def normalize_label(
    value
):

    value = normalize_text(
        value
    ).lower()

    value = value.replace(
        "&",
        " and ",
    )

    value = re.sub(
        r"\s+",
        " ",
        value,
    )

    return value.strip()


def numeric(
    value
):

    if value is None:

        return np.nan


    if isinstance(
        value,
        (
            int,
            float,
            np.integer,
            np.floating,
        ),
    ):

        return float(value)


    text = normalize_text(
        value
    )


    if text in {
        "",
        "-",
        "—",
        "–",
    }:

        return np.nan


    negative = (
        text.startswith("(")
        and
        text.endswith(")")
    )


    text = (
        text
        .replace("$", "")
        .replace(",", "")
        .replace("%", "")
        .replace("(", "")
        .replace(")", "")
        .strip()
    )


    try:

        result = float(
            text
        )

    except Exception:

        return np.nan


    if negative:

        result = -abs(
            result
        )


    return result


# ============================================================
# 4. GROWTH FUNCTIONS
# ============================================================

def log_growth(
    series,
    periods=4,
):

    series = pd.to_numeric(
        series,
        errors="coerce",
    )


    lagged = (
        series.shift(
            periods
        )
    )


    output = pd.Series(
        np.nan,
        index=series.index,
        dtype=float,
    )


    valid = (
        (series > 0)
        &
        (lagged > 0)
    )


    output.loc[
        valid
    ] = (
        np.log(
            series.loc[
                valid
            ]
            /
            lagged.loc[
                valid
            ]
        )
        * 100
    )


    return output


def log_to_normal_pct(
    value
):

    if isinstance(
        value,
        pd.Series,
    ):

        return (
            np.exp(
                value / 100
            )
            - 1
        ) * 100


    if isinstance(
        value,
        np.ndarray,
    ):

        return (
            np.exp(
                value / 100
            )
            - 1
        ) * 100


    if pd.isna(value):

        return np.nan


    return (
        np.exp(
            value / 100
        )
        - 1
    ) * 100


def mae(
    actual,
    predicted,
):

    actual = np.asarray(
        actual,
        dtype=float,
    )


    predicted = np.asarray(
        predicted,
        dtype=float,
    )


    return float(
        np.mean(
            np.abs(
                actual
                -
                predicted
            )
        )
    )


def mae_yoy_pct_points(
    actual_log,
    predicted_log,
):

    actual_pct = log_to_normal_pct(
        np.asarray(
            actual_log,
            dtype=float,
        )
    )


    predicted_pct = log_to_normal_pct(
        np.asarray(
            predicted_log,
            dtype=float,
        )
    )


    return float(
        np.mean(
            np.abs(
                actual_pct
                -
                predicted_pct
            )
        )
    )


# ============================================================
# 5. LOAD BASE FILES
# ============================================================

PANEL_FILE = lake_path(
    "energy_v2_1_panel.csv"
)


VALIDATION_FILE = lake_path(
    "energy_v2_1_validation.csv"
)


NOWCAST_FILE = lake_path(
    "energy_v2_1_nowcast.csv"
)


for file_path in [
    PANEL_FILE,
    VALIDATION_FILE,
    NOWCAST_FILE,
]:

    if not file_path.exists():

        raise FileNotFoundError(
            f"{file_path} 없음."
        )


# ------------------------------------------------------------
# Structural base
#
# champion V3.2.3 우선
# ------------------------------------------------------------

STRUCTURAL_CANDIDATES = [

    lake_path(
        "energy_v3_2_3_structural_panel.csv"
    ),

    lake_path(
        "energy_v3_2_5_structural_panel.csv"
    ),

    lake_path(
        "energy_v3_2_4_structural_panel.csv"
    ),
]


STRUCTURAL_BASE_FILE = None


for candidate in (
    STRUCTURAL_CANDIDATES
):

    if candidate.exists():

        STRUCTURAL_BASE_FILE = (
            candidate
        )

        break


if STRUCTURAL_BASE_FILE is None:

    raise FileNotFoundError(
        "V3.2.3/3.2.4/3.2.5 "
        "structural panel이 없어."
    )


print(
    "\nStructural base:",
    STRUCTURAL_BASE_FILE
)


# ============================================================
# 6. LOAD DATA
# ============================================================

panel_v21 = pd.read_csv(
    PANEL_FILE
)


panel_v21[
    "quarter"
] = pd.PeriodIndex(
    panel_v21[
        "quarter"
    ].astype(str),
    freq="Q",
)


validation_v21 = pd.read_csv(
    VALIDATION_FILE
)


validation_v21[
    "quarter"
] = pd.PeriodIndex(
    validation_v21[
        "quarter"
    ].astype(str),
    freq="Q",
)


nowcast_v21 = pd.read_csv(
    NOWCAST_FILE
)


nowcast_v21[
    "nowcast_quarter"
] = pd.PeriodIndex(
    nowcast_v21[
        "nowcast_quarter"
    ].astype(str),
    freq="Q",
)


base_structural = pd.read_csv(
    STRUCTURAL_BASE_FILE
)


base_structural[
    "quarter"
] = pd.PeriodIndex(
    base_structural[
        "quarter"
    ].astype(str),
    freq="Q",
)


base_structural = (
    base_structural[
        base_structural[
            "ticker"
        ].isin(
            TICKERS
        )
    ]
    .copy()
)


# ============================================================
# 7. EOG EVENT QUARTER
# ============================================================

QUARTER_WORD = {

    "first":
        1,

    "second":
        2,

    "third":
        3,

    "fourth":
        4,
}


def parse_event_quarter(
    text
):

    text = normalize_text(
        text
    )


    match = re.search(
        r"EOG\s+Resources\s+"
        r"(First|Second|Third|Fourth)\s+Quarter"
        r"(?:\s+and\s+Full[-\s]?Year)?"
        r"\s+(20\d{2})"
        r"\s+Webcast\s+and\s+Results",
        text,
        flags=re.I,
    )


    if not match:

        return None


    q = QUARTER_WORD[
        match.group(1).lower()
    ]


    year = int(
        match.group(2)
    )


    return pd.Period(
        f"{year}Q{q}",
        freq="Q",
    )


# ============================================================
# 8. PARSE EVENT DETAIL
# ============================================================

def parse_event_detail(
    url,
):

    try:

        response = http_get(
            url
        )

    except Exception:

        return None


    soup = BeautifulSoup(
        response.text,
        "lxml",
    )


    page_text = normalize_text(
        soup.get_text(
            " ",
            strip=True,
        )
    )


    result_q = parse_event_quarter(
        page_text
    )


    if result_q is None:

        return None


    stats_url = None
    guidance_url = None
    press_release_url = None


    for anchor in soup.find_all(
        "a",
        href=True,
    ):

        anchor_text = normalize_label(
            anchor.get_text(
                " ",
                strip=True,
            )
        )


        href = urljoin(
            EOG_BASE_URL,
            anchor[
                "href"
            ],
        )


        if (
            "statistics"
            in anchor_text
            and
            "xls"
            in anchor_text
        ):

            stats_url = href


        elif (
            anchor_text
            ==
            "guidance"
        ):

            guidance_url = href


        elif (
            anchor_text
            ==
            "press release"
        ):

            press_release_url = href


    return {

        "quarter":
            result_q,

        "event_url":
            url,

        "stats_url":
            stats_url,

        "guidance_url":
            guidance_url,

        "press_release_url":
            press_release_url,
    }


# ============================================================
# 9. DISCOVER EOG EARNINGS EVENTS
# ============================================================

def discover_eog_events():

    candidate_urls = set()


    # --------------------------------------------------------
    # Main event listing
    # --------------------------------------------------------

    for listing_url in [
        EOG_EVENTS_URL,
        EOG_INVESTORS_URL,
    ]:

        try:

            response = http_get(
                listing_url
            )


            soup = BeautifulSoup(
                response.text,
                "lxml",
            )


            for anchor in soup.find_all(
                "a",
                href=True,
            ):

                href = anchor[
                    "href"
                ]


                if (
                    "events-and-presentations"
                    in href
                    and
                    "item="
                    in href
                ):

                    candidate_urls.add(
                        urljoin(
                            EOG_BASE_URL,
                            href,
                        )
                    )


        except Exception:

            pass


    details = {}


    # --------------------------------------------------------
    # Parse discovered links
    # --------------------------------------------------------

    for url in sorted(
        candidate_urls
    ):

        detail = parse_event_detail(
            url
        )


        if (
            detail is None
            or
            detail[
                "quarter"
            ].year
            <
            EOG_IR_START_YEAR
        ):

            continue


        details[
            detail[
                "quarter"
            ]
        ] = detail


    # --------------------------------------------------------
    # Fallback:
    #
    # item IDs 직접 검사
    # --------------------------------------------------------

    if (
        len(details)
        <
        EOG_MIN_EVENTS
    ):

        visited = set(
            candidate_urls
        )


        for item_id in range(
            EOG_EVENT_ITEM_MAX,
            0,
            -1,
        ):

            url = (
                EOG_EVENTS_URL
                +
                f"?item={item_id}"
            )


            if url in visited:

                continue


            detail = parse_event_detail(
                url
            )


            time.sleep(
                0.03
            )


            if detail is None:

                continue


            if (
                detail[
                    "quarter"
                ].year
                <
                EOG_IR_START_YEAR
            ):

                continue


            details[
                detail[
                    "quarter"
                ]
            ] = detail


            if (
                len(details)
                >=
                EOG_MIN_EVENTS
                and
                min(
                    x.year
                    for x
                    in details
                )
                <=
                EOG_IR_START_YEAR
            ):

                break


    result = sorted(
        details.values(),
        key=lambda x:
            x[
                "quarter"
            ],
    )


    event_df = pd.DataFrame(
        result
    )


    event_df.to_csv(
        lake_path(
            "energy_v3_2_6_eog_events.csv"
        ),
        index=False,
    )


    print(
        "\nEOG events discovered:",
        len(
            result
        )
    )


    if not event_df.empty:

        print(
            event_df[
                [
                    "quarter",
                    "stats_url",
                    "guidance_url",
                ]
            ]
            .tail(12)
            .to_string(
                index=False
            )
        )


    return result


# ============================================================
# 10. XLS HEADER PARSING
# ============================================================

def detect_quarter_token(
    value
):

    text = normalize_label(
        value
    )


    patterns = [

        (
            r"\b([1-4])"
            r"(?:st|nd|rd|th)"
            r"\s*qtr\b",
            1,
        ),

        (
            r"\bq([1-4])\b",
            1,
        ),

        (
            r"\b([1-4])q\b",
            1,
        ),
    ]


    for pattern, group in patterns:

        match = re.search(
            pattern,
            text,
        )


        if match:

            return int(
                match.group(
                    group
                )
            )


    return None


def detect_year(
    value
):

    text = normalize_text(
        value
    )


    match = re.search(
        r"\b(20\d{2})\b",
        text,
    )


    if match:

        return int(
            match.group(1)
        )


    return None


def value_at(
    data,
    row,
    col,
):

    if (
        row < 0
        or
        row >= len(data)
    ):

        return None


    if (
        col < 0
        or
        col >= len(
            data[row]
        )
    ):

        return None


    return data[
        row
    ][
        col
    ]


def infer_period_for_column(
    data,
    row_index,
    column_index,
):

    start_row = max(
        0,
        row_index
        -
        15,
    )


    quarter_number = None
    year = None


    # --------------------------------------------------------
    # Same-column quarter/year
    # --------------------------------------------------------

    for r in range(
        start_row,
        row_index,
    ):

        cell = value_at(
            data,
            r,
            column_index,
        )


        q = detect_quarter_token(
            cell
        )


        y = detect_year(
            cell
        )


        if q is not None:

            quarter_number = q


        if y is not None:

            year = y


    # --------------------------------------------------------
    # Year merged cell may be to the left
    # --------------------------------------------------------

    if (
        quarter_number
        is not None
        and
        year
        is None
    ):

        for distance in range(
            0,
            10,
        ):

            c = (
                column_index
                -
                distance
            )


            if c < 0:

                break


            for r in range(
                start_row,
                row_index,
            ):

                y = detect_year(
                    value_at(
                        data,
                        r,
                        c,
                    )
                )


                if y is not None:

                    year = y

                    break


            if year is not None:

                break


    if (
        quarter_number
        is None
        or
        year
        is None
    ):

        return None


    return pd.Period(
        f"{year}Q{quarter_number}",
        freq="Q",
    )


# ============================================================
# 11. XLS COMPONENT LABEL
# ============================================================

def classify_xls_section(
    label
):

    label = normalize_label(
        label
    )


    if (
        "average"
        in label
        and
        "price"
        in label
    ):

        return None


    if (
        "crude oil and condensate"
        in label
        and
        "volume"
        in label
    ):

        return "oil"


    if (
        "natural gas liquids"
        in label
        and
        "volume"
        in label
    ):

        return "ngl"


    if (
        "natural gas"
        in label
        and
        "liquid"
        not in label
        and
        "volume"
        in label
    ):

        return "gas"


    if (
        "crude oil equivalent"
        in label
        and
        "volume"
        in label
    ):

        return "total"


    return None


# ============================================================
# 12. XLS TARGET VALUE
# ============================================================

def find_quarter_value(
    data,
    row_index,
    start_column,
    target_q,
):

    max_columns = max(
        len(row)
        for row
        in data
    )


    for c in range(
        start_column,
        max_columns,
    ):

        period = (
            infer_period_for_column(
                data,
                row_index,
                c,
            )
        )


        if period != target_q:

            continue


        value = numeric(
            value_at(
                data,
                row_index,
                c,
            )
        )


        if pd.notna(
            value
        ):

            return float(
                value
            )


    return np.nan


# ============================================================
# 13. PARSE EOG STATISTICS XLS
# ============================================================

def parse_eog_statistics_xls(
    url,
    target_q,
):

    response = http_get(
        url
    )


    workbook = load_workbook(
        BytesIO(
            response.content
        ),
        read_only=True,
        data_only=True,
    )


    candidates = {

        "oil":
            [],

        "ngl":
            [],

        "gas":
            [],

        "total":
            [],
    }


    debug_rows = []


    for worksheet in (
        workbook.worksheets
    ):

        data = [
            list(row)

            for row
            in worksheet.iter_rows(
                values_only=True
            )
        ]


        if not data:

            continue


        section = None


        for r, row in enumerate(
            data
        ):

            first_col = None
            label = ""


            for c, cell in enumerate(
                row
            ):

                if (
                    cell is not None
                    and
                    normalize_text(
                        cell
                    )
                ):

                    first_col = c

                    label = (
                        normalize_text(
                            cell
                        )
                    )

                    break


            if first_col is None:

                continue


            component = (
                classify_xls_section(
                    label
                )
            )


            if (
                component
                is not None
            ):

                section = (
                    component
                )


                # Direct row value
                direct_value = (
                    find_quarter_value(
                        data,
                        r,
                        first_col
                        +
                        1,
                        target_q,
                    )
                )


                if pd.notna(
                    direct_value
                ):

                    candidates[
                        component
                    ].append(
                        {
                            "value":
                                direct_value,

                            "sheet":
                                worksheet.title,

                            "row":
                                r + 1,

                            "label":
                                label,

                            "type":
                                "DIRECT",
                        }
                    )


                continue


            # ------------------------------------------------
            # Section Total row
            # ------------------------------------------------

            label_norm = normalize_label(
                label
            )


            if (
                section
                is not None
                and
                re.fullmatch(
                    r"total",
                    label_norm,
                )
            ):

                value = find_quarter_value(
                    data,
                    r,
                    first_col
                    +
                    1,
                    target_q,
                )


                if pd.notna(
                    value
                ):

                    candidates[
                        section
                    ].append(
                        {
                            "value":
                                value,

                            "sheet":
                                worksheet.title,

                            "row":
                                r + 1,

                            "label":
                                label,

                            "type":
                                "TOTAL_ROW",
                        }
                    )


        # debug candidate rows
        for component, rows in (
            candidates.items()
        ):

            for item in rows:

                debug_rows.append(
                    {
                        "quarter":
                            str(
                                target_q
                            ),

                        "component":
                            component,

                        **item,
                    }
                )


    # --------------------------------------------------------
    # Pick best
    # --------------------------------------------------------

    selected = {}


    for component, items in (
        candidates.items()
    ):

        if not items:

            selected[
                component
            ] = np.nan

            continue


        # Total-row 우선,
        # direct fallback
        total_rows = [
            x
            for x
            in items
            if x[
                "type"
            ]
            ==
            "TOTAL_ROW"
        ]


        pick_from = (
            total_rows
            if total_rows
            else items
        )


        selected[
            component
        ] = float(
            pick_from[
                -1
            ][
                "value"
            ]
        )


    return (
        selected,
        debug_rows,
    )


# ============================================================
# 14. BOE RECONSTRUCTION
# ============================================================

def reconstruct_total_boe(
    oil,
    ngl,
    gas,
):

    if (
        pd.isna(oil)
        or
        pd.isna(ngl)
        or
        pd.isna(gas)
    ):

        return np.nan


    return float(
        oil
        +
        ngl
        +
        gas
        /
        6.0
    )


def reconstruction_error_pct(
    reported,
    reconstructed,
):

    if (
        pd.isna(
            reported
        )
        or
        pd.isna(
            reconstructed
        )
        or
        reported
        <= 0
    ):

        return np.nan


    return float(
        abs(
            reported
            -
            reconstructed
        )
        /
        reported
        *
        100
    )


def finalize_eog_actual(
    quarter,
    values,
    source_url,
):

    oil = values.get(
        "oil",
        np.nan,
    )

    ngl = values.get(
        "ngl",
        np.nan,
    )

    gas = values.get(
        "gas",
        np.nan,
    )

    reported = values.get(
        "total",
        np.nan,
    )


    reconstructed = (
        reconstruct_total_boe(
            oil,
            ngl,
            gas,
        )
    )


    error = (
        reconstruction_error_pct(
            reported,
            reconstructed,
        )
    )


    if (
        pd.notna(
            reported
        )
        and
        pd.notna(
            reconstructed
        )
    ):

        if (
            error
            <=
            EOG_RECON_ERROR_MAX_PCT
        ):

            total = (
                reported
            )

            qc_flag = (
                "REPORTED_RECON_MATCH"
            )

            quality = 1.00


        else:

            total = (
                reconstructed
            )

            qc_flag = (
                "USE_RECONSTRUCTED"
            )

            quality = 0.85


    elif pd.notna(
        reconstructed
    ):

        total = (
            reconstructed
        )

        qc_flag = (
            "RECONSTRUCTED_ONLY"
        )

        quality = 0.90


    elif pd.notna(
        reported
    ):

        total = (
            reported
        )

        qc_flag = (
            "REPORTED_ONLY"
        )

        quality = 0.70


    else:

        total = np.nan

        qc_flag = (
            "NO_VALID_TOTAL"
        )

        quality = 0.0


    return {

        "quarter":
            quarter,

        "oil_production":
            oil,

        "ngl_production":
            ngl,

        "gas_production":
            gas,

        "reported_total_production":
            reported,

        "reconstructed_total_production":
            reconstructed,

        "reconstruction_error_pct":
            error,

        "total_production":
            total,

        "component_qc_flag":
            qc_flag,

        "data_quality_score":
            quality,

        "source_url":
            source_url,
    }


# ============================================================
# 15. GUIDANCE PDF TEXT
# ============================================================

def pdf_text_from_url(
    url,
):

    response = http_get(
        url
    )


    try:

        reader = PdfReader(
            BytesIO(
                response.content
            )
        )


        text = "\n".join(
            (
                page.extract_text()
                or
                ""
            )

            for page
            in reader.pages
        )


        return normalize_text(
            text
        )


    except Exception:

        # HTML/text fallback
        try:

            return normalize_text(
                BeautifulSoup(
                    response.text,
                    "lxml",
                ).get_text(
                    " ",
                    strip=True,
                )
            )

        except Exception:

            return ""


# ============================================================
# 16. GUIDANCE TARGET QUARTER
# ============================================================

def infer_guidance_quarter(
    text,
):

    lower = normalize_label(
        text
    )


    # 3Q 2026 Guidance Range
    match = re.search(
        r"\b([1-4])q\s*(20\d{2})"
        r".{0,80}?"
        r"guidance",
        lower,
    )


    if match:

        return pd.Period(
            (
                f"{match.group(2)}"
                f"Q{match.group(1)}"
            ),
            freq="Q",
        )


    # Third Quarter ... 2026 Guidance
    match = re.search(
        r"(first|second|third|fourth)"
        r"\s+quarter"
        r".{0,100}?"
        r"(20\d{2})"
        r".{0,60}?"
        r"guidance",
        lower,
    )


    if match:

        quarter_number = (
            QUARTER_WORD[
                match.group(1)
            ]
        )


        return pd.Period(
            (
                f"{match.group(2)}"
                f"Q{quarter_number}"
            ),
            freq="Q",
        )


    return None


# ============================================================
# 17. GUIDANCE NUMBER EXTRACTION
# ============================================================

def extract_numbers_from_text(
    text,
):

    tokens = re.findall(
        r"\(?-?"
        r"\d[\d,]*"
        r"(?:\.\d+)?"
        r"\)?",
        text,
    )


    result = []


    for token in tokens:

        value = numeric(
            token
        )


        if pd.isna(value):

            continue


        if (
            1990
            <= value
            <= 2100
        ):

            continue


        result.append(
            float(value)
        )


    return result


def midpoint_from_values(
    values,
    low_limit,
    high_limit,
):

    values = [
        x
        for x in values
        if (
            low_limit
            <= x
            <= high_limit
        )
    ]


    if len(
        values
    ) >= 3:

        calculated_midpoint = (
            values[0]
            +
            values[1]
        ) / 2


        if abs(
            values[2]
            -
            calculated_midpoint
        ) <= max(
            1.0,
            abs(
                calculated_midpoint
            )
            *
            0.04,
        ):

            return float(
                values[2]
            )


    if len(
        values
    ) >= 2:

        return float(
            (
                values[0]
                +
                values[1]
            )
            /
            2
        )


    if len(
        values
    ) == 1:

        return float(
            values[0]
        )


    return np.nan


# ============================================================
# 18. GUIDANCE SECTION
# ============================================================

def extract_guidance_section_total(
    text,
    section_pattern,
    next_patterns,
    low_limit,
    high_limit,
):

    match = re.search(
        section_pattern,
        text,
        flags=re.I,
    )


    if not match:

        return np.nan


    start = (
        match.end()
    )


    end = min(
        len(text),
        start
        +
        1800,
    )


    for next_pattern in (
        next_patterns
    ):

        next_match = re.search(
            next_pattern,
            text[
                start:
            ],
            flags=re.I,
        )


        if next_match:

            possible_end = (
                start
                +
                next_match.start()
            )


            if possible_end > start:

                end = min(
                    end,
                    possible_end,
                )


    section = text[
        start:
        end
    ]


    total_matches = list(
        re.finditer(
            r"\bTotal\b",
            section,
            flags=re.I,
        )
    )


    for total_match in (
        total_matches
    ):

        tail = section[
            total_match.end():
            total_match.end()
            +
            300
        ]


        numbers = (
            extract_numbers_from_text(
                tail
            )
        )


        midpoint = midpoint_from_values(
            numbers,
            low_limit,
            high_limit,
        )


        if pd.notna(
            midpoint
        ):

            return midpoint


    return np.nan


# ============================================================
# 19. PARSE EOG GUIDANCE
# ============================================================

def parse_eog_guidance(
    url,
):

    text = pdf_text_from_url(
        url
    )


    if not text:

        return None


    target_q = (
        infer_guidance_quarter(
            text
        )
    )


    if target_q is None:

        return None


    oil = extract_guidance_section_total(

        text,

        (
            r"Crude\s+Oil\s+and\s+"
            r"Condensate\s+Volumes?"
            r"\s*\(MB(?:o|bl)d\)"
        ),

        [
            r"Natural\s+Gas\s+Liquids",
        ],

        10,
        2000,
    )


    ngl = extract_guidance_section_total(

        text,

        (
            r"Natural\s+Gas\s+Liquids"
            r"\s+Volumes?"
            r"\s*\(MBbld\)"
        ),

        [
            r"Natural\s+Gas\s+Volumes",
        ],

        0,
        2000,
    )


    gas = extract_guidance_section_total(

        text,

        (
            r"Natural\s+Gas\s+Volumes?"
            r"\s*\(MMcfd\)"
        ),

        [
            r"Crude\s+Oil\s+Equivalent",
        ],

        0,
        20000,
    )


    reported_total = (
        extract_guidance_section_total(

            text,

            (
                r"(?:Total\s+)?"
                r"Crude\s+Oil\s+Equivalent"
                r"\s+Volumes?"
                r"\s*\(MBo(?:e)?d\)"
            ),

            [
                r"Crude\s+Oil\s+and\s+"
                r"Condensate\s*-\s*above",
                r"Benchmark\s+Price",
            ],

            100,
            4000,
        )
    )


    reconstructed = (
        reconstruct_total_boe(
            oil,
            ngl,
            gas,
        )
    )


    error = (
        reconstruction_error_pct(
            reported_total,
            reconstructed,
        )
    )


    if (
        pd.notna(
            reported_total
        )
        and
        pd.notna(
            reconstructed
        )
    ):

        if (
            error
            <=
            EOG_RECON_ERROR_MAX_PCT
        ):

            final_total = (
                reported_total
            )

            flag = (
                "REPORTED_RECON_MATCH"
            )

            quality = 1.00


        else:

            final_total = (
                reconstructed
            )

            flag = (
                "USE_RECONSTRUCTED"
            )

            quality = 0.85


    elif pd.notna(
        reconstructed
    ):

        final_total = (
            reconstructed
        )

        flag = (
            "RECONSTRUCTED_ONLY"
        )

        quality = 0.90


    elif pd.notna(
        reported_total
    ):

        final_total = (
            reported_total
        )

        flag = (
            "REPORTED_ONLY"
        )

        quality = 0.70


    else:

        return None


    return {

        "target_quarter":
            target_q,

        "guidance_oil_production":
            oil,

        "guidance_ngl_production":
            ngl,

        "guidance_gas_production":
            gas,

        "guidance_reported_total":
            reported_total,

        "guidance_reconstructed_total":
            reconstructed,

        "guidance_reconstruction_error_pct":
            error,

        "guidance_total_production":
            final_total,

        "guidance_component_qc_flag":
            flag,

        "guidance_data_quality_score":
            quality,

        "source_url":
            url,
    }


# ============================================================
# 20. BUILD EOG IR CACHE
# ============================================================

EOG_ACTUAL_CACHE = lake_path(
    "energy_v3_2_6_actual_EOG.csv"
)


EOG_GUIDANCE_CACHE = lake_path(
    "energy_v3_2_6_guidance_EOG.csv"
)


EOG_XLS_DEBUG_CACHE = lake_path(
    "energy_v3_2_6_eog_xls_debug.csv"
)


if (
    REFRESH_EOG_IR
    or
    not EOG_ACTUAL_CACHE.exists()
    or
    not EOG_GUIDANCE_CACHE.exists()
):

    events = discover_eog_events()


    actual_rows = []

    guidance_rows = []

    xls_debug = []


    for event in events:

        result_q = event[
            "quarter"
        ]


        if (
            result_q.year
            <
            EOG_IR_START_YEAR
        ):

            continue


        # ----------------------------------------------------
        # Actual XLS
        # ----------------------------------------------------

        if event[
            "stats_url"
        ]:

            try:

                (
                    xls_values,
                    debug_rows,
                ) = (
                    parse_eog_statistics_xls(
                        event[
                            "stats_url"
                        ],
                        result_q,
                    )
                )


                actual_row = (
                    finalize_eog_actual(
                        result_q,
                        xls_values,
                        event[
                            "stats_url"
                        ],
                    )
                )


                if pd.notna(
                    actual_row[
                        "total_production"
                    ]
                ):

                    actual_rows.append(
                        actual_row
                    )


                xls_debug.extend(
                    debug_rows
                )


            except Exception as exc:

                print(
                    f"⚠️ EOG {result_q} XLS 실패:",
                    repr(exc),
                )


        # ----------------------------------------------------
        # Guidance
        # ----------------------------------------------------

        if event[
            "guidance_url"
        ]:

            try:

                guidance_row = (
                    parse_eog_guidance(
                        event[
                            "guidance_url"
                        ]
                    )
                )


                if (
                    guidance_row
                    is not None
                ):

                    guidance_rows.append(
                        guidance_row
                    )


            except Exception as exc:

                print(
                    f"⚠️ EOG {result_q} guidance 실패:",
                    repr(exc),
                )


    # --------------------------------------------------------
    # DataFrames
    # --------------------------------------------------------

    eog_actual = pd.DataFrame(
        actual_rows
    )


    eog_guidance = pd.DataFrame(
        guidance_rows
    )


    # --------------------------------------------------------
    # Deduplicate
    # --------------------------------------------------------

    if not eog_actual.empty:

        eog_actual[
            "quarter"
        ] = pd.PeriodIndex(
            eog_actual[
                "quarter"
            ].astype(str),
            freq="Q",
        )


        eog_actual = (
            eog_actual
            .sort_values(
                [
                    "quarter",
                    "data_quality_score",
                ]
            )
            .drop_duplicates(
                "quarter",
                keep="last",
            )
            .sort_values(
                "quarter"
            )
        )


    if not eog_guidance.empty:

        eog_guidance[
            "target_quarter"
        ] = pd.PeriodIndex(
            eog_guidance[
                "target_quarter"
            ].astype(str),
            freq="Q",
        )


        eog_guidance = (
            eog_guidance
            .sort_values(
                [
                    "target_quarter",
                    "guidance_data_quality_score",
                ]
            )
            .drop_duplicates(
                "target_quarter",
                keep="last",
            )
            .sort_values(
                "target_quarter"
            )
        )


    eog_actual.to_csv(
        EOG_ACTUAL_CACHE,
        index=False,
    )


    eog_guidance.to_csv(
        EOG_GUIDANCE_CACHE,
        index=False,
    )


    pd.DataFrame(
        xls_debug
    ).to_csv(
        EOG_XLS_DEBUG_CACHE,
        index=False,
    )


else:

    eog_actual = pd.read_csv(
        EOG_ACTUAL_CACHE
    )


    eog_guidance = pd.read_csv(
        EOG_GUIDANCE_CACHE
    )


    if not eog_actual.empty:

        eog_actual[
            "quarter"
        ] = pd.PeriodIndex(
            eog_actual[
                "quarter"
            ].astype(str),
            freq="Q",
        )


    if not eog_guidance.empty:

        eog_guidance[
            "target_quarter"
        ] = pd.PeriodIndex(
            eog_guidance[
                "target_quarter"
            ].astype(str),
            freq="Q",
        )


# ============================================================
# 21. EOG IR CHECK
# ============================================================

print(
    "\n"
    + "=" * 90
)

print(
    "EOG IR XLS / GUIDANCE CHECK"
)

print(
    "=" * 90
)


if not eog_actual.empty:

    print(
        eog_actual[
            [
                "quarter",

                "oil_production",

                "ngl_production",

                "gas_production",

                "reported_total_production",

                "reconstructed_total_production",

                "reconstruction_error_pct",

                "total_production",

                "component_qc_flag",

                "data_quality_score",
            ]
        ]
        .tail(16)
        .round(2)
        .to_string(
            index=False
        )
    )


if not eog_guidance.empty:

    print(
        "\n===== GUIDANCE ====="
    )


    print(
        eog_guidance[
            [
                "target_quarter",

                "guidance_oil_production",

                "guidance_ngl_production",

                "guidance_gas_production",

                "guidance_reported_total",

                "guidance_reconstructed_total",

                "guidance_reconstruction_error_pct",

                "guidance_total_production",

                "guidance_component_qc_flag",

                "guidance_data_quality_score",
            ]
        ]
        .tail(10)
        .round(2)
        .to_string(
            index=False
        )
    )


# ============================================================
# 22. EOG CURRENT DATA CHECK
# ============================================================

eog_last_revenue_q = (
    panel_v21[
        (
            panel_v21[
                "ticker"
            ]
            ==
            "EOG"
        )
        &
        (
            panel_v21[
                "revenue"
            ].notna()
        )
    ][
        "quarter"
    ]
    .max()
)


eog_nowcast_q = (
    eog_last_revenue_q
    +
    1
)


latest_ir_actual_q = (
    eog_actual[
        "quarter"
    ].max()

    if not eog_actual.empty

    else None
)


guidance_quarters = (
    set(
        eog_guidance[
            "target_quarter"
        ]
    )

    if not eog_guidance.empty

    else set()
)


print(
    "\nEOG latest revenue quarter :",
    eog_last_revenue_q
)

print(
    "EOG latest IR actual      :",
    latest_ir_actual_q
)

print(
    "EOG target nowcast        :",
    eog_nowcast_q
)

print(
    "EOG target guidance found :",
    (
        eog_nowcast_q
        in guidance_quarters
    )
)


# ============================================================
# 23. PERIOD MAPPER
# ============================================================

def map_period_values(
    quarter_series,
    source_df,
    index_column,
    value_column,
    default=np.nan,
):

    if (
        source_df is None
        or
        source_df.empty
        or
        value_column
        not in source_df.columns
    ):

        return pd.Series(
            default,
            index=quarter_series.index,
        )


    mapping = dict(
        zip(
            source_df[
                index_column
            ],
            source_df[
                value_column
            ],
        )
    )


    return pd.Series(
        [
            mapping.get(
                q,
                default,
            )

            for q in quarter_series
        ],
        index=quarter_series.index,
    )


# ============================================================
# 24. BUILD EOG STRUCTURAL ROWS
# ============================================================

def rebuild_eog_structural(
    base_eog,
    actual_df,
    guidance_df,
):

    df = (
        base_eog
        .sort_values(
            "quarter"
        )
        .reset_index(
            drop=True
        )
        .copy()
    )


    # ========================================================
    # Actual production
    # ========================================================

    for source_col in [

        "total_production",

        "oil_production",

        "ngl_production",

        "gas_production",

        "data_quality_score",

    ]:

        target_col = (

            "actual_data_quality"

            if source_col
            ==
            "data_quality_score"

            else source_col

        )


        df[
            target_col
        ] = map_period_values(

            df[
                "quarter"
            ],

            actual_df,

            "quarter",

            source_col,

            default=(
                0.0
                if target_col
                ==
                "actual_data_quality"
                else
                np.nan
            ),

        )


    # ========================================================
    # Guidance
    # ========================================================

    df[
        "guidance_total_production"
    ] = map_period_values(

        df[
            "quarter"
        ],

        guidance_df,

        "target_quarter",

        "guidance_total_production",

    )


    df[
        "guidance_data_quality"
    ] = map_period_values(

        df[
            "quarter"
        ],

        guidance_df,

        "target_quarter",

        "guidance_data_quality_score",

        default=0.0,
    )


    # ========================================================
    # Actual production YoY
    # ========================================================

    df[
        "total_prod_yoy"
    ] = log_growth(
        df[
            "total_production"
        ],
        4,
    )


    df[
        "total_prod_yoy_l1"
    ] = (
        df[
            "total_prod_yoy"
        ]
        .shift(1)
    )


    # Pair quality
    pair_quality = (
        pd.concat(
            [
                df[
                    "actual_data_quality"
                ],

                df[
                    "actual_data_quality"
                ]
                .shift(4),
            ],
            axis=1,
        )
        .min(
            axis=1
        )
    )


    df[
        "actual_pair_quality_l1"
    ] = (
        pair_quality
        .shift(1)
    )


    # ========================================================
    # Guidance vs Actual
    # ========================================================

    prior_actual = (
        df[
            "total_production"
        ]
        .shift(4)
    )


    prior_guidance = (
        df[
            "guidance_total_production"
        ]
        .shift(4)
    )


    df[
        "guidance_vs_actual_yoy"
    ] = np.nan


    valid = (
        (
            df[
                "guidance_total_production"
            ]
            > 0
        )
        &
        (
            prior_actual
            > 0
        )
    )


    df.loc[
        valid,
        "guidance_vs_actual_yoy"
    ] = (
        np.log(
            df.loc[
                valid,
                "guidance_total_production"
            ]
            /
            prior_actual.loc[
                valid
            ]
        )
        * 100
    )


    guidance_actual_quality = (
        pd.concat(
            [
                df[
                    "guidance_data_quality"
                ],

                df[
                    "actual_data_quality"
                ]
                .shift(4),
            ],
            axis=1,
        )
        .min(
            axis=1
        )
    )


    # ========================================================
    # Guidance vs Guidance
    # ========================================================

    df[
        "guidance_vs_guidance_yoy"
    ] = np.nan


    valid = (
        (
            df[
                "guidance_total_production"
            ]
            > 0
        )
        &
        (
            prior_guidance
            > 0
        )
    )


    df.loc[
        valid,
        "guidance_vs_guidance_yoy"
    ] = (
        np.log(
            df.loc[
                valid,
                "guidance_total_production"
            ]
            /
            prior_guidance.loc[
                valid
            ]
        )
        * 100
    )


    guidance_guidance_quality = (
        pd.concat(
            [
                df[
                    "guidance_data_quality"
                ],

                df[
                    "guidance_data_quality"
                ]
                .shift(4),
            ],
            axis=1,
        )
        .min(
            axis=1
        )
    )


    # ========================================================
    # Industry fallback
    # ========================================================

    industry_prod = (
        0.70
        *
        df[
            "crude_prod_yoy_l1"
        ]
        +
        0.30
        *
        df[
            "gas_prod_yoy_l1"
        ]
    )


    # ========================================================
    # Production source
    # ========================================================

    production_proxy = []

    production_source = []

    quality_score = []


    for i in range(
        len(df)
    ):

        # ----------------------------------------------------
        # Guidance vs actual
        # ----------------------------------------------------

        if pd.notna(
            df.loc[
                i,
                "guidance_vs_actual_yoy"
            ]
        ):

            value = df.loc[
                i,
                "guidance_vs_actual_yoy"
            ]


            source = (
                "GUIDANCE_VS_ACTUAL"
            )


            quality = (
                SOURCE_QUALITY[
                    source
                ]
                *
                guidance_actual_quality.iloc[
                    i
                ]
            )


        # ----------------------------------------------------
        # Guidance vs guidance
        # ----------------------------------------------------

        elif pd.notna(
            df.loc[
                i,
                "guidance_vs_guidance_yoy"
            ]
        ):

            value = df.loc[
                i,
                "guidance_vs_guidance_yoy"
            ]


            source = (
                "GUIDANCE_VS_GUIDANCE"
            )


            quality = (
                SOURCE_QUALITY[
                    source
                ]
                *
                guidance_guidance_quality.iloc[
                    i
                ]
            )


        # ----------------------------------------------------
        # Actual vs actual
        # ----------------------------------------------------

        elif pd.notna(
            df.loc[
                i,
                "total_prod_yoy_l1"
            ]
        ):

            value = df.loc[
                i,
                "total_prod_yoy_l1"
            ]


            source = (
                "ACTUAL_VS_ACTUAL"
            )


            pair_quality_value = (
                df.loc[
                    i,
                    "actual_pair_quality_l1"
                ]
            )


            if pd.isna(
                pair_quality_value
            ):

                pair_quality_value = 0.0


            quality = (
                SOURCE_QUALITY[
                    source
                ]
                *
                pair_quality_value
            )


        # ----------------------------------------------------
        # Industry
        # ----------------------------------------------------

        else:

            value = (
                industry_prod.iloc[
                    i
                ]
            )


            source = (
                "INDUSTRY_FALLBACK"
            )


            quality = (
                SOURCE_QUALITY[
                    source
                ]
            )


        if pd.notna(
            value
        ):

            value = float(
                np.clip(
                    value,
                    -60,
                    120,
                )
            )


        production_proxy.append(
            value
        )


        production_source.append(
            source
        )


        quality_score.append(
            float(
                np.clip(
                    quality,
                    0,
                    1,
                )
            )
        )


    df[
        "company_prod_yoy_proxy"
    ] = production_proxy


    df[
        "production_yoy_source"
    ] = production_source


    df[
        "source_quality_score"
    ] = quality_score


    df[
        "has_company_prod"
    ] = (
        df[
            "production_yoy_source"
        ]
        .ne(
            "INDUSTRY_FALLBACK"
        )
        .astype(float)
    )


    df[
        "has_guidance"
    ] = (
        df[
            "guidance_total_production"
        ]
        .notna()
        .astype(float)
    )


    # ========================================================
    # Oil share
    # ========================================================

    oil_share = (
        df[
            "oil_production"
        ]
        /
        df[
            "total_production"
        ]
    )


    oil_share = (
        oil_share
        .replace(
            [
                np.inf,
                -np.inf,
            ],
            np.nan,
        )
        .clip(
            0.10,
            1.00,
        )
    )


    # EOG 실제 recent mix fallback
    median_oil_share = (
        oil_share.median(
            skipna=True
        )
    )


    if pd.isna(
        median_oil_share
    ):

        median_oil_share = 0.39


    df[
        "oil_share_l1"
    ] = (
        oil_share
        .shift(1)
        .fillna(
            median_oil_share
        )
        .clip(
            0.10,
            1.00,
        )
    )


    # ========================================================
    # Price mix
    #
    # NGL은 oil-like로 처리
    # non-oil BOE 절반만 Henry Hub 직접 노출
    # ========================================================

    gas_weight = (
        (
            1.0
            -
            df[
                "oil_share_l1"
            ]
        )
        *
        0.50
    )


    oil_like_weight = (
        1.0
        -
        gas_weight
    )


    df[
        "price_mix_yoy"
    ] = (
        oil_like_weight
        *
        df[
            "wti_yoy"
        ]
        +
        gas_weight
        *
        df[
            "henry_yoy"
        ]
    )


    # ========================================================
    # Structural Revenue
    # ========================================================

    df[
        "structural_revenue_log_yoy"
    ] = (
        df[
            "price_mix_yoy"
        ]
        +
        df[
            "company_prod_yoy_proxy"
        ]
    )


    df[
        "structural_revenue_log_yoy"
    ] = (
        df[
            "structural_revenue_log_yoy"
        ]
        .clip(
            MIN_LOG_GROWTH,
            MAX_LOG_GROWTH,
        )
    )


    df[
        "structural_revenue_yoy_pct"
    ] = log_to_normal_pct(
        df[
            "structural_revenue_log_yoy"
        ]
    )


    return df


# ============================================================
# 25. PATCH EOG ONLY
# ============================================================

non_eog = (
    base_structural[
        base_structural[
            "ticker"
        ]
        !=
        "EOG"
    ]
    .copy()
)


base_eog = (
    base_structural[
        base_structural[
            "ticker"
        ]
        ==
        "EOG"
    ]
    .copy()
)


patched_eog = (
    rebuild_eog_structural(
        base_eog,
        eog_actual,
        eog_guidance,
    )
)


structural_panel = (
    pd.concat(
        [
            non_eog,
            patched_eog,
        ],
        ignore_index=True,
        sort=False,
    )
    .sort_values(
        [
            "quarter",
            "ticker",
        ]
    )
    .reset_index(
        drop=True
    )
)


structural_panel.to_csv(
    lake_path(
        "energy_v3_2_6_structural_panel.csv"
    ),
    index=False,
)


# ============================================================
# 26. SOURCE SUMMARY
# ============================================================

source_summary = (
    structural_panel
    .groupby(
        [
            "ticker",
            "production_yoy_source",
        ]
    )
    .agg(
        observations=(
            "production_yoy_source",
            "size",
        ),

        avg_quality=(
            "source_quality_score",
            "mean",
        ),
    )
    .reset_index()
)


source_summary.to_csv(
    lake_path(
        "energy_v3_2_6_source_quality_summary.csv"
    ),
    index=False,
)


print(
    "\n"
    + "=" * 90
)

print(
    "PRODUCTION SOURCE SUMMARY"
)

print(
    "=" * 90
)


print(
    source_summary
    .round(2)
    .to_string(
        index=False
    )
)


# ============================================================
# 27. VALIDATION ALIGNMENT
# ============================================================

structural_validation = (
    structural_panel[
        structural_panel[
            "revenue_yoy"
        ]
        .notna()
    ][
        [
            "quarter",

            "ticker",

            "revenue_yoy",

            "structural_revenue_log_yoy",

            "company_prod_yoy_proxy",

            "production_yoy_source",

            "source_quality_score",

            "has_company_prod",
        ]
    ]
    .copy()
    .rename(
        columns={
            "revenue_yoy":
                "actual_log_yoy"
        }
    )
)


v21_validation = (
    validation_v21[
        [
            "quarter",

            "ticker",

            "actual",

            "predicted",
        ]
    ]
    .copy()
    .rename(
        columns={
            "actual":
                "v21_actual_log_yoy",

            "predicted":
                "v21_log_yoy",
        }
    )
)


compare = (
    v21_validation.merge(
        structural_validation,
        on=[
            "quarter",
            "ticker",
        ],
        how="inner",
    )
)


compare[
    "actual_log_yoy"
] = (
    compare[
        "v21_actual_log_yoy"
    ]
)


# ============================================================
# 28. BLEND HELPERS
# ============================================================

def best_weight(
    history
):

    if history.empty:

        return 0.0


    best_w = 0.0
    best_error = np.inf


    for weight in (
        BLEND_WEIGHTS
    ):

        prediction = (
            weight
            *
            history[
                "structural_revenue_log_yoy"
            ]
            +
            (
                1.0
                -
                weight
            )
            *
            history[
                "v21_log_yoy"
            ]
        )


        error = mae(
            history[
                "actual_log_yoy"
            ],
            prediction,
        )


        if (
            error
            <
            best_error
        ):

            best_error = (
                error
            )

            best_w = float(
                weight
            )


    return best_w


def guarded_base_weight(
    history
):

    if history.empty:

        return 0.0


    structural_error = mae(
        history[
            "actual_log_yoy"
        ],
        history[
            "structural_revenue_log_yoy"
        ],
    )


    v21_error = mae(
        history[
            "actual_log_yoy"
        ],
        history[
            "v21_log_yoy"
        ],
    )


    if (
        structural_error
        >=
        v21_error
    ):

        return 0.0


    return best_weight(
        history
    )


def quality_adjusted_weight(
    base_weight,
    row
):

    quality = float(
        row[
            "source_quality_score"
        ]
    )


    effective = (
        base_weight
        *
        quality
    )


    if (
        row[
            "production_yoy_source"
        ]
        ==
        "INDUSTRY_FALLBACK"
    ):

        effective = min(
            effective,
            NO_COMPANY_PROD_MAX_WEIGHT,
        )


    return float(
        np.clip(
            effective,
            0,
            1,
        )
    )


# ============================================================
# 29. WALK-FORWARD
# ============================================================

walk_rows = []


for ticker, company_df in (
    compare.groupby(
        "ticker"
    )
):

    company_df = (
        company_df
        .sort_values(
            "quarter"
        )
        .reset_index(
            drop=True
        )
    )


    for i in range(
        len(company_df)
    ):

        row = company_df.iloc[
            i
        ]


        history = (
            company_df.iloc[
                :i
            ]
        )


        if (
            len(history)
            <
            MIN_BLEND_HISTORY
        ):

            base_weight = 0.0

        else:

            base_weight = (
                guarded_base_weight(
                    history
                )
            )


        effective_weight = (
            quality_adjusted_weight(
                base_weight,
                row,
            )
        )


        prediction = (
            effective_weight
            *
            row[
                "structural_revenue_log_yoy"
            ]
            +
            (
                1.0
                -
                effective_weight
            )
            *
            row[
                "v21_log_yoy"
            ]
        )


        prediction = float(
            np.clip(
                prediction,
                MIN_LOG_GROWTH,
                MAX_LOG_GROWTH,
            )
        )


        walk_rows.append(
            {
                "quarter":
                    row[
                        "quarter"
                    ],

                "ticker":
                    ticker,

                "actual_log_yoy":
                    row[
                        "actual_log_yoy"
                    ],

                "actual_yoy_pct":
                    log_to_normal_pct(
                        row[
                            "actual_log_yoy"
                        ]
                    ),

                "v21_log_yoy":
                    row[
                        "v21_log_yoy"
                    ],

                "v21_yoy_pct":
                    log_to_normal_pct(
                        row[
                            "v21_log_yoy"
                        ]
                    ),

                "structural_log_yoy":
                    row[
                        "structural_revenue_log_yoy"
                    ],

                "structural_yoy_pct":
                    log_to_normal_pct(
                        row[
                            "structural_revenue_log_yoy"
                        ]
                    ),

                "base_structural_weight":
                    base_weight,

                "source_quality_score":
                    row[
                        "source_quality_score"
                    ],

                "effective_structural_weight":
                    effective_weight,

                "blend_log_yoy":
                    prediction,

                "blend_yoy_pct":
                    log_to_normal_pct(
                        prediction
                    ),

                "company_prod_yoy_proxy":
                    row[
                        "company_prod_yoy_proxy"
                    ],

                "production_yoy_source":
                    row[
                        "production_yoy_source"
                    ],
            }
        )


walk_validation = pd.DataFrame(
    walk_rows
)


# ============================================================
# 30. METRICS
# ============================================================

metric_rows = []


print(
    "\n"
    + "=" * 90
)

print(
    "V3.2.6 WALK-FORWARD VALIDATION"
)

print(
    "=" * 90
)


for ticker, group in (
    walk_validation.groupby(
        "ticker"
    )
):

    metric_rows.append(
        {
            "ticker":
                ticker,

            "V2_1_MAE_log_points":
                mae(
                    group[
                        "actual_log_yoy"
                    ],
                    group[
                        "v21_log_yoy"
                    ],
                ),

            "V2_1_MAE_yoy_pct_points":
                mae_yoy_pct_points(
                    group[
                        "actual_log_yoy"
                    ],
                    group[
                        "v21_log_yoy"
                    ],
                ),

            "Structural_MAE_log_points":
                mae(
                    group[
                        "actual_log_yoy"
                    ],
                    group[
                        "structural_log_yoy"
                    ],
                ),

            "Structural_MAE_yoy_pct_points":
                mae_yoy_pct_points(
                    group[
                        "actual_log_yoy"
                    ],
                    group[
                        "structural_log_yoy"
                    ],
                ),

            "V3_2_6_MAE_log_points":
                mae(
                    group[
                        "actual_log_yoy"
                    ],
                    group[
                        "blend_log_yoy"
                    ],
                ),

            "V3_2_6_MAE_yoy_pct_points":
                mae_yoy_pct_points(
                    group[
                        "actual_log_yoy"
                    ],
                    group[
                        "blend_log_yoy"
                    ],
                ),
        }
    )


metrics = (
    pd.DataFrame(
        metric_rows
    )
    .set_index(
        "ticker"
    )
)


overall_v21_log = mae(
    walk_validation[
        "actual_log_yoy"
    ],
    walk_validation[
        "v21_log_yoy"
    ],
)


overall_v21_pct = (
    mae_yoy_pct_points(
        walk_validation[
            "actual_log_yoy"
        ],
        walk_validation[
            "v21_log_yoy"
        ],
    )
)


overall_structural_log = mae(
    walk_validation[
        "actual_log_yoy"
    ],
    walk_validation[
        "structural_log_yoy"
    ],
)


overall_structural_pct = (
    mae_yoy_pct_points(
        walk_validation[
            "actual_log_yoy"
        ],
        walk_validation[
            "structural_log_yoy"
        ],
    )
)


overall_v326_log = mae(
    walk_validation[
        "actual_log_yoy"
    ],
    walk_validation[
        "blend_log_yoy"
    ],
)


overall_v326_pct = (
    mae_yoy_pct_points(
        walk_validation[
            "actual_log_yoy"
        ],
        walk_validation[
            "blend_log_yoy"
        ],
    )
)


print(
    "\n===== LOG-GROWTH SCALE ====="
)

print(
    f"V2.1       : "
    f"{overall_v21_log:.2f} log-points"
)

print(
    f"Structural : "
    f"{overall_structural_log:.2f} log-points"
)

print(
    f"V3.2.6     : "
    f"{overall_v326_log:.2f} log-points"
)


print(
    "\n===== NORMAL REVENUE YOY SCALE ====="
)

print(
    f"V2.1       : "
    f"{overall_v21_pct:.2f} %p"
)

print(
    f"Structural : "
    f"{overall_structural_pct:.2f} %p"
)

print(
    f"V3.2.6     : "
    f"{overall_v326_pct:.2f} %p"
)


print(
    "\n===== BY TICKER ====="
)

print(
    metrics
    .round(2)
)


walk_validation.to_csv(
    lake_path(
        "energy_v3_2_6_validation.csv"
    ),
    index=False,
)


metrics.to_csv(
    lake_path(
        "energy_v3_2_6_metrics.csv"
    )
)


# ============================================================
# 31. FINAL BASE WEIGHTS
# ============================================================

final_base_weights = {}

weight_rows = []


for ticker, group in (
    compare.groupby(
        "ticker"
    )
):

    structural_error = mae(
        group[
            "actual_log_yoy"
        ],
        group[
            "structural_revenue_log_yoy"
        ],
    )


    v21_error = mae(
        group[
            "actual_log_yoy"
        ],
        group[
            "v21_log_yoy"
        ],
    )


    if (
        structural_error
        >=
        v21_error
    ):

        weight = 0.0

        guardrail = (
            "STRUCTURAL_DISABLED"
        )

    else:

        weight = best_weight(
            group
        )

        guardrail = (
            "STRUCTURAL_ENABLED"
        )


    final_base_weights[
        ticker
    ] = weight


    weight_rows.append(
        {
            "ticker":
                ticker,

            "base_structural_weight":
                weight,

            "historical_structural_MAE_log_points":
                structural_error,

            "historical_v21_MAE_log_points":
                v21_error,

            "guardrail":
                guardrail,
        }
    )


weights_df = (
    pd.DataFrame(
        weight_rows
    )
    .set_index(
        "ticker"
    )
)


print(
    "\n===== FINAL BASE WEIGHTS ====="
)

print(
    weights_df
    .round(3)
)


weights_df.to_csv(
    lake_path(
        "energy_v3_2_6_blend_weights.csv"
    )
)


# ============================================================
# 32. CURRENT STRUCTURAL ROWS
# ============================================================

structural_nowcast_rows = []


for ticker in TICKERS:

    company = (
        structural_panel[
            structural_panel[
                "ticker"
            ]
            ==
            ticker
        ]
        .copy()
    )


    actual_revenue = (
        company[
            company[
                "revenue"
            ]
            .notna()
        ]
    )


    if actual_revenue.empty:

        continue


    last_q = (
        actual_revenue[
            "quarter"
        ]
        .max()
    )


    target_q = (
        last_q
        +
        1
    )


    target = (
        company[
            company[
                "quarter"
            ]
            ==
            target_q
        ]
    )


    if target.empty:

        continue


    row = target.iloc[
        0
    ]


    structural_nowcast_rows.append(
        {
            "ticker":
                ticker,

            "nowcast_quarter":
                target_q,

            "structural_log_yoy":
                row[
                    "structural_revenue_log_yoy"
                ],

            "structural_yoy_pct":
                row[
                    "structural_revenue_yoy_pct"
                ],

            "price_mix_yoy":
                row[
                    "price_mix_yoy"
                ],

            "company_prod_yoy_proxy":
                row[
                    "company_prod_yoy_proxy"
                ],

            "production_yoy_source":
                row[
                    "production_yoy_source"
                ],

            "source_quality_score":
                row[
                    "source_quality_score"
                ],

            "guidance_total_production":
                (
                    row[
                        "guidance_total_production"
                    ]
                    if
                    "guidance_total_production"
                    in row.index
                    else
                    np.nan
                ),

            "has_company_prod":
                row[
                    "has_company_prod"
                ],
        }
    )


structural_nowcast = pd.DataFrame(
    structural_nowcast_rows
)


structural_nowcast[
    "nowcast_quarter"
] = pd.PeriodIndex(
    structural_nowcast[
        "nowcast_quarter"
    ].astype(str),
    freq="Q",
)


# ============================================================
# 33. MERGE CURRENT V2.1
# ============================================================

v21_nowcast = (
    nowcast_v21[
        [
            "ticker",

            "nowcast_quarter",

            "predicted_revenue_yoy",
        ]
    ]
    .copy()
    .rename(
        columns={
            "predicted_revenue_yoy":
                "v21_log_yoy"
        }
    )
)


current = (
    structural_nowcast.merge(
        v21_nowcast,
        on=[
            "ticker",
            "nowcast_quarter",
        ],
        how="inner",
    )
)


# ============================================================
# 34. CURRENT WEIGHTS
# ============================================================

base_weights = []

effective_weights = []


for _, row in (
    current.iterrows()
):

    base = (
        final_base_weights.get(
            row[
                "ticker"
            ],
            0.0,
        )
    )


    effective = (
        quality_adjusted_weight(
            base,
            row,
        )
    )


    base_weights.append(
        base
    )


    effective_weights.append(
        effective
    )


current[
    "base_structural_weight"
] = base_weights


current[
    "structural_weight"
] = effective_weights


current[
    "v21_weight"
] = (
    1.0
    -
    current[
        "structural_weight"
    ]
)


# ============================================================
# 35. FINAL NOWCAST
# ============================================================

current[
    "predicted_revenue_log_yoy"
] = (
    current[
        "structural_weight"
    ]
    *
    current[
        "structural_log_yoy"
    ]
    +
    current[
        "v21_weight"
    ]
    *
    current[
        "v21_log_yoy"
    ]
)


current[
    "predicted_revenue_log_yoy"
] = (
    current[
        "predicted_revenue_log_yoy"
    ]
    .clip(
        MIN_LOG_GROWTH,
        MAX_LOG_GROWTH,
    )
)


current[
    "predicted_revenue_yoy_pct"
] = (
    log_to_normal_pct(
        current[
            "predicted_revenue_log_yoy"
        ]
    )
)


current[
    "v21_yoy_pct"
] = (
    log_to_normal_pct(
        current[
            "v21_log_yoy"
        ]
    )
)


# ============================================================
# 36. REVENUE LEVEL
# ============================================================

predicted_revenues = []


for _, row in (
    current.iterrows()
):

    base_q = (
        row[
            "nowcast_quarter"
        ]
        -
        4
    )


    base = (
        panel_v21[
            (
                panel_v21[
                    "ticker"
                ]
                ==
                row[
                    "ticker"
                ]
            )
            &
            (
                panel_v21[
                    "quarter"
                ]
                ==
                base_q
            )
        ]
    )


    if (
        base.empty
        or
        pd.isna(
            base.iloc[
                0
            ][
                "revenue"
            ]
        )
    ):

        predicted_revenues.append(
            np.nan
        )

        continue


    base_revenue = float(
        base.iloc[
            0
        ][
            "revenue"
        ]
    )


    revenue = (
        base_revenue
        *
        np.exp(
            row[
                "predicted_revenue_log_yoy"
            ]
            /
            100
        )
    )


    predicted_revenues.append(
        revenue
    )


current[
    "predicted_revenue"
] = predicted_revenues


current[
    "predicted_revenue_B"
] = (
    current[
        "predicted_revenue"
    ]
    /
    1e9
)


# ============================================================
# 37. CONFIDENCE
# ============================================================

current[
    "model_spread_log_points"
] = (
    current[
        "structural_log_yoy"
    ]
    -
    current[
        "v21_log_yoy"
    ]
).abs()


def selected_model(
    row
):

    if (
        row[
            "structural_weight"
        ]
        <=
        0.001
    ):

        return (
            "V2.1_ONLY"
        )


    if (
        row[
            "structural_weight"
        ]
        >=
        0.999
    ):

        return (
            "STRUCTURAL_ONLY"
        )


    return "BLEND"


def confidence(
    row
):

    if (
        row[
            "base_structural_weight"
        ]
        <= 0
    ):

        return "LOW"


    if (
        row[
            "production_yoy_source"
        ]
        ==
        "INDUSTRY_FALLBACK"
    ):

        return "LOW"


    if (
        row[
            "source_quality_score"
        ]
        >=
        0.90
        and
        row[
            "model_spread_log_points"
        ]
        <=
        10
    ):

        return "HIGH"


    if (
        row[
            "source_quality_score"
        ]
        >=
        0.70
        and
        row[
            "model_spread_log_points"
        ]
        <=
        20
    ):

        return "MEDIUM"


    return "LOW"


current[
    "selected_model"
] = current.apply(
    selected_model,
    axis=1,
)


current[
    "confidence"
] = current.apply(
    confidence,
    axis=1,
)


# ============================================================
# 38. FINAL OUTPUT
# ============================================================

output = (
    current[
        [
            "ticker",

            "nowcast_quarter",

            "predicted_revenue_B",

            "predicted_revenue_log_yoy",

            "predicted_revenue_yoy_pct",

            "v21_log_yoy",

            "v21_yoy_pct",

            "structural_log_yoy",

            "structural_yoy_pct",

            "base_structural_weight",

            "source_quality_score",

            "structural_weight",

            "v21_weight",

            "selected_model",

            "company_prod_yoy_proxy",

            "production_yoy_source",

            "guidance_total_production",

            "model_spread_log_points",

            "confidence",
        ]
    ]
    .copy()
)


print(
    "\n"
    + "=" * 90
)

print(
    "🚀 ENERGY REVENUE V3.2.6 NOWCAST"
)

print(
    "=" * 90
)


print(
    output
    .round(2)
    .to_string(
        index=False
    )
)


output.to_csv(
    lake_path(
        "energy_v3_2_6_nowcast.csv"
    ),
    index=False,
)


# ============================================================
# 39. COMPARE V3.2.5
# ============================================================

old_metrics_file = lake_path(
    "energy_v3_2_5_metrics.csv"
)


if old_metrics_file.exists():

    try:

        old_metrics = pd.read_csv(
            old_metrics_file,
            index_col=0,
        )


        print(
            "\n"
            + "=" * 90
        )

        print(
            "V3.2.5 vs V3.2.6"
        )

        print(
            "=" * 90
        )


        comparison = metrics.copy()


        if (
            "V3_2_5_MAE_log_points"
            in old_metrics.columns
        ):

            comparison[
                "V3_2_5_log_MAE"
            ] = (
                old_metrics[
                    "V3_2_5_MAE_log_points"
                ]
            )


            comparison[
                "log_improvement"
            ] = (
                comparison[
                    "V3_2_5_log_MAE"
                ]
                -
                comparison[
                    "V3_2_6_MAE_log_points"
                ]
            )


        if (
            "V3_2_5_MAE_yoy_pct_points"
            in old_metrics.columns
        ):

            comparison[
                "V3_2_5_pct_MAE"
            ] = (
                old_metrics[
                    "V3_2_5_MAE_yoy_pct_points"
                ]
            )


            comparison[
                "pct_improvement"
            ] = (
                comparison[
                    "V3_2_5_pct_MAE"
                ]
                -
                comparison[
                    "V3_2_6_MAE_yoy_pct_points"
                ]
            )


        print(
            comparison
            .round(2)
        )


    except Exception as exc:

        print(
            "Comparison skipped:",
            repr(exc),
        )


# ============================================================
# 40. METADATA
# ============================================================

metadata = {

    "version":
        "3.2.6",

    "structural_base_file":
        str(
            STRUCTURAL_BASE_FILE
        ),

    "eog_data_source":
        (
            "Official EOG Investor Relations "
            "Statistics XLS + Guidance PDF"
        ),

    "eog_ir_start_year":
        EOG_IR_START_YEAR,

    "eog_reconstruction_formula":
        (
            "Oil MBod + NGL MBbld "
            "+ Gas MMcfd / 6 "
            "= Total MBoed"
        ),

    "eog_max_reconstruction_error_pct":
        EOG_RECON_ERROR_MAX_PCT,

    "production_source_priority": [

        "GUIDANCE_VS_ACTUAL",

        "GUIDANCE_VS_GUIDANCE",

        "ACTUAL_VS_ACTUAL",

        "INDUSTRY_FALLBACK",
    ],

    "metrics": [

        "MAE_log_points",

        "MAE_yoy_pct_points",
    ],
}


with open(
    lake_path(
        "energy_v3_2_6_metadata.json"
    ),
    "w",
    encoding="utf-8",
) as f:

    json.dump(
        metadata,
        f,
        indent=2,
        ensure_ascii=False,
    )


# ============================================================
# 41. DONE
# ============================================================

print(
    "\n"
    + "=" * 90
)

print(
    "DONE"
)

print(
    "=" * 90
)


print(
    "\nMain outputs:"
)


for filename in [

    "energy_v3_2_6_eog_events.csv",

    "energy_v3_2_6_actual_EOG.csv",

    "energy_v3_2_6_guidance_EOG.csv",

    "energy_v3_2_6_eog_xls_debug.csv",

    "energy_v3_2_6_structural_panel.csv",

    "energy_v3_2_6_source_quality_summary.csv",

    "energy_v3_2_6_validation.csv",

    "energy_v3_2_6_metrics.csv",

    "energy_v3_2_6_blend_weights.csv",

    "energy_v3_2_6_nowcast.csv",

    "energy_v3_2_6_metadata.json",

]:

    print(
        " -",
        lake_path(
            filename
        )
    )
